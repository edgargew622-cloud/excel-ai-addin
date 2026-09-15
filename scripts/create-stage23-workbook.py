from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName

root = Path(__file__).resolve().parent.parent
target = root / "stage2-3-smoke.xlsx"

wb = Workbook()
ws = wb.active
ws.title = "Продажи"
rows = [
    ["Дата", "Город", "Товар", "Количество", "Цена", "Сумма", "Старая ошибка", "Результат агента"],
    ["2026-09-01", "Москва", "Кофе", 2, 450, "=D2*E2", "=1/0", None],
    ["2026-09-02", "Казань", "Чай", 3, 280, "=D3*E3", None, None],
    ["2026-09-03", "Москва", "Какао", 1, 520, "=D4*E4", None, None],
    ["2026-09-04", "Омск", "Кофе", 4, 450, "=D5*E5", None, None],
    ["2026-09-05", "Казань", "Чай", 2, 280, "=D6*E6", None, None],
]
for row in rows:
    ws.append(row)
for cell in ws[1]:
    cell.font = Font(bold=True)

table = Table(displayName="SalesTable", ref="A1:H6")
table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(table)
wb.defined_names.add(DefinedName("SalesBlock", attr_text="'Продажи'!$A$1:$F$6"))

chart = BarChart()
chart.title = "Сумма по строкам"
chart.add_data(Reference(ws, min_col=6, min_row=1, max_row=6), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=3, min_row=2, max_row=6))
ws.add_chart(chart, "J2")

lookup = wb.create_sheet("Справочник")
lookup.append(["Товар", "Категория"])
lookup.append(["Кофе", "Напитки"])
lookup.append(["Чай", "Напитки"])
lookup.append(["Какао", "Напитки"])

wb.create_sheet("Пустой")
hidden = wb.create_sheet("Скрытый")
hidden["A1"] = "Этот лист скрыт"
hidden.sheet_state = "hidden"

protected = wb.create_sheet("Защищённый")
protected["A1"] = "Запись запрещена защитой листа"
protected.protection.sheet = True

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.save(target)
print(target)
