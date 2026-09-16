"""Эталонная книга для проверок, как её описывает раздел 10 плана V4.

Прежние книги делались под конкретный этап и покрывали не всё: в них не было
ни формул между листами, ни диаграммы рядом с защитой. Здесь собрано всё, что
план требует от тестовой книги, чтобы проверки оформления, структуры, сортировки
и резервных копий шли на одном материале.

Чего здесь нет и почему: сводную таблицу openpyxl не создаёт — её добавляют
руками в Excel, когда дойдёт черёд до create_pivot_table. Отдельный файл
другого формата создаётся рядом как .csv.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo

root = Path(__file__).resolve().parent.parent
target = root / "reference-smoke.xlsx"
csv_target = root / "reference-smoke.csv"

wb = Workbook()

# --- Продажи: таблица, формулы, заранее существующая ошибка -------------------
sales = wb.active
sales.title = "Продажи"
sales.append(["Дата", "Город", "Товар", "Количество", "Цена", "Сумма", "Ошибка"])
rows = [
    ("2026-09-01", "Москва", "Кофе", 2, 450),
    ("2026-09-02", "Казань", "Чай", 3, 280),
    ("2026-09-03", "Москва", "Какао", 1, 520),
    ("2026-09-04", "Омск", "Кофе", 4, 450),
    ("2026-09-05", "Казань", "Чай", 2, 280),
]
for index, (date, city, item, count, price) in enumerate(rows, start=2):
    sales.append([date, city, item, count, price, f"=D{index}*E{index}", None])
# Ошибка существует до любых наших правок: по ней проверяется, что агент
# не выдаёт чужие ошибки за последствия своей записи.
sales["G2"] = "=1/0"

for cell in sales[1]:
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor="DDEBF7")
    cell.alignment = Alignment(horizontal="center")

table = Table(displayName="SalesTable", ref="A1:G6")
table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
sales.add_table(table)

chart = BarChart()
chart.title = "Сумма по строкам"
chart.add_data(Reference(sales, min_col=6, min_row=1, max_row=6), titles_from_data=True)
chart.set_categories(Reference(sales, min_col=3, min_row=2, max_row=6))
sales.add_chart(chart, "I2")

# Объединение: по нему проверяется предупреждение перед записью.
sales.merge_cells("A9:C9")
sales["A9"] = "Объединено A9:C9"

# --- Справочник ---------------------------------------------------------------
lookup = wb.create_sheet("Справочник")
lookup.append(["Товар", "Категория", "Наценка"])
for item, group, markup in [("Кофе", "Напитки", 0.15), ("Чай", "Напитки", 0.10), ("Какао", "Напитки", 0.20)]:
    lookup.append([item, group, markup])

# --- Свод: формулы между листами ----------------------------------------------
summary = wb.create_sheet("Свод")
summary.append(["Город", "Выручка", "Доля"])
for index, city in enumerate(["Москва", "Казань", "Омск"], start=2):
    summary.append([
        city,
        f"=SUMIF(Продажи!$B$2:$B$6,A{index},Продажи!$F$2:$F$6)",
        f"=B{index}/SUM($B$2:$B$4)",
    ])
summary["E1"] = "Всего товаров"
summary["E2"] = "=COUNTA(Справочник!A2:A4)"

# --- Защищённый и пустой -------------------------------------------------------
protected = wb.create_sheet("Защищённый")
protected["A1"] = "Запись запрещена защитой листа"
protected["A2"] = "=Продажи!F2"
protected.protection.sheet = True

wb.create_sheet("Пустой")

hidden = wb.create_sheet("Скрытый")
hidden["A1"] = "Этот лист скрыт"
hidden.sheet_state = "hidden"

wb.defined_names.add(DefinedName("SalesBlock", attr_text="'Продажи'!$A$1:$F$6"))
wb.calculation.fullCalcOnLoad = True
wb.save(target)

# Отдельный файл другого формата: план требует проверить, что копия не
# переименовывает чужой формат в .xlsx.
csv_target.write_text(
    "Город;Выручка\nМосква;1420\nКазань;1400\nОмск;1800\n",
    encoding="utf-8-sig",
)

print(f"Создано: {target}")
print(f"Создано: {csv_target}")
print("Сводную таблицу добавьте вручную в Excel, когда дойдёт черёд до create_pivot_table.")
