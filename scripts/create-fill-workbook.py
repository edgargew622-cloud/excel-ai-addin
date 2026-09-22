"""Книга для проверки заполнения формулами: fill-smoke.xlsx (план стабилизации, S1).

Прежний запасной путь заполнения портил ссылки, похожие на адрес ячейки:
`=Q1!A1` становилось `=Q2!A2`, `=SUM(Sales[Q1])` — `=SUM(Sales[Q2])`.
Здесь такие имена стоят нарочно, и числа на листах разные, поэтому порча
видна не только по формуле, но и по значению:

    лист Q1        A1:A5 = 10, 20, 30, 40, 50;   B1:C1 = 11, 12
    лист Q2        A1:A5 = 1, 2, 3, 4, 5
    лист Q1 2026   A1:A5 = 100, 200, 300, 400, 500

На листе «Данные» — таблица Sales (A1:C6) со столбцами Q1 и Q2:
сумма Q1 = 150, сумма Q2 = 15. Столбцы E и правее пусты: туда заполняем,
не вплотную к таблице, чтобы она не расширялась.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo

root = Path(__file__).resolve().parent.parent
target = root / "fill-smoke.xlsx"

wb = Workbook()
data = wb.active
data.title = "Данные"
data.append(["Товар", "Q1", "Q2"])
for index, (q1, q2) in enumerate(zip([10, 20, 30, 40, 50], [1, 2, 3, 4, 5]), start=1):
    data.append([f"Товар {index}", q1, q2])
table = Table(displayName="Sales", ref="A1:C6")
table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
data.add_table(table)
for cell in data[1]:
    cell.font = Font(bold=True)

q1 = wb.create_sheet("Q1")
for row, value in enumerate([10, 20, 30, 40, 50], start=1):
    q1.cell(row, 1, value)
q1["B1"] = 11
q1["C1"] = 12

q2 = wb.create_sheet("Q2")
for row, value in enumerate([1, 2, 3, 4, 5], start=1):
    q2.cell(row, 1, value)

q2026 = wb.create_sheet("Q1 2026")
for row, value in enumerate([100, 200, 300, 400, 500], start=1):
    q2026.cell(row, 1, value)

wb.save(target)
print(f"Книга сохранена: {target}")
