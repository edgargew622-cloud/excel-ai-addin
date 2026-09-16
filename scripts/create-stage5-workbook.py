"""Книга для проверок этапа 5: групповая запись и замер выгрузки.

Отдельная от stage2-3-smoke.xlsx, чтобы проверки не зависели от накопленных
в той книге ручных правок. Свободные столбцы G:J оставлены под записи,
объединение L1:N1 нужно для проверки предупреждения, защищённый лист —
для проверки остановки группы на сбое.
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font

root = Path(__file__).resolve().parent.parent
target = root / "stage5-group-smoke.xlsx"

wb = Workbook()
ws = wb.active
ws.title = "Данные"
ws.append(["Товар", "Цена", "Количество", "Сумма"])
for cell in ws[1]:
    cell.font = Font(bold=True)
for row in [["Кофе", 450, 2], ["Чай", 280, 3], ["Какао", 520, 1], ["Сок", 190, 4], ["Вода", 60, 6]]:
    ws.append(row + [f"=B{ws.max_row + 1}*C{ws.max_row + 1}"])

# Объединение для проверки предупреждения о записи внутрь объединённой области.
ws.merge_cells("L1:N1")
ws["L1"] = "Объединено L1:N1"

protected = wb.create_sheet("Защищённый")
protected["A1"] = "Запись запрещена защитой листа"
protected.protection.sheet = True

wb.create_sheet("Пустой")

wb.calculation.fullCalcOnLoad = True
wb.save(target)
print(f"Создано: {target}")
